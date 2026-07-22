# -*- coding: utf-8 -*-
"""
build_navy_catalog.py — navy 스킨 카탈로그 (세로 6장)
판형 8.27 x 11.69in (A4 세로)

구성: [표지] [철학·강점] [커리큘럼] [실적·특별프로그램] [시간표·학습관리] [입학·FAQ·연락처]
데이터 없는 섹션은 건너뛰고, 페이지가 비면 그 장을 만들지 않는다.

입력: v3 표준스키마 (PPT·리플렛과 동일)
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

from build_brandbook_navy import _style_table, _set_cell
from navy_core import (
    FS, get_palette, blank_slide, add_box, add_text, add_pill,
    place_image, clean, fit_size, gv, as_dicts, wrap_words,
)

SW, SH = 8.27, 11.69
ML = 0.63
CW = SW - ML * 2          # 7.01
R = 0.055                 # 공통 라운드 반경


def head(s, P, eyebrow, title, lead=None, y=0.72):
    add_text(s, ML, y, CW, 0.2, eyebrow.upper(), size=FS.eyebrow, bold=True,
             color=P["primary"], line_spacing=1.0)
    lines = wrap_words(clean(title), 22)[:2]
    size = 19.5 if len(lines) == 1 else 17.25
    h = len(lines) * size * 1.28 / 72.0 + 0.06
    add_text(s, ML, y + 0.28, CW, h, lines, size=size, bold=True,
             color=P["text"], line_spacing=1.28)
    yy = y + 0.28 + h
    if lead:
        ll = wrap_words(clean(lead), 42)[:2]
        lh = len(ll) * FS.body * 1.4 / 72.0 + 0.04
        add_text(s, ML, yy + 0.12, CW, lh, ll, size=FS.body,
                 color=P["muted"], line_spacing=1.4)
        yy += 0.12 + lh
    return yy + 0.34


def footer(s, P, d, page):
    b = d.get("basic", {})
    add_box(s, ML, SH - 0.92, CW, 0.01, fill=P["line"])
    add_text(s, ML, SH - 0.74, CW * 0.7, 0.22, clean(b.get("name") or ""),
             size=FS.small, bold=True, color=P["muted"], line_spacing=1.0)
    add_text(s, SW - ML - 0.6, SH - 0.74, 0.6, 0.22, f"{page:02d}",
             size=FS.tiny, bold=True, color=P["muted"],
             align=PP_ALIGN.RIGHT, line_spacing=1.0)


# ── 1. 표지 ─────────────────────────────────────────────────────
def c_cover(prs, d, P):
    s = blank_slide(prs)
    b = d.get("basic", {})
    add_box(s, 0, 0, SW, SH, fill=P["deep"])

    photo = gv(d, "assets", "cover") or gv(d, "assets", "photo")
    if photo:
        place_image(s, photo, 0, 0, SW, SH * 0.52, cover=True)
        add_box(s, 0, SH * 0.44, SW, SH * 0.08, fill=P["deep"])

    y = SH * 0.55 if photo else 2.4
    logo = gv(d, "assets", "logo")
    if logo:
        place_image(s, logo, ML, y, 2.2, 1.1, cover=False)
        y += 1.45

    target = clean(b.get("target") or "")
    if target:
        tw = min(CW, 0.55 + len(target) * 0.15)
        add_pill(s, ML, y, tw, 0.32, target, fill=P["accent"],
                 color=P["deep"], size=FS.body)
        y += 0.78

    slogan = clean(b.get("slogan") or gv(d, "identity", "slogan") or "")
    if slogan:
        lines = wrap_words(slogan, 15)[:3]
        add_text(s, ML, y, CW, len(lines) * 26 * 1.32 / 72.0 + 0.1, lines,
                 size=26, bold=True, color=P["onDark"], line_spacing=1.32,
                 clip=False)
        y += len(lines) * 26 * 1.32 / 72.0 + 0.4

    sub = clean(b.get("subline") or "")
    if sub:
        add_text(s, ML, y, CW, 0.5, wrap_words(sub, 34)[:2], size=FS.lead,
                 color=P["onDarkSub"], line_spacing=1.45)

    add_text(s, ML, SH - 1.55, CW, 0.34, clean(b.get("name") or ""),
             size=FS.card_ttl, bold=True, color=P["onDark"], line_spacing=1.0)
    foot = "  |  ".join([x for x in (clean(b.get("phone") or ""),
                                     clean(b.get("address") or "")) if x])
    if foot:
        add_text(s, ML, SH - 1.12, CW, 0.28, foot, size=FS.small,
                 color=P["onDarkSub"], line_spacing=1.0)
    return s


# ── 2. 철학 + 강점 ──────────────────────────────────────────────
def c_philosophy(prs, d, P, page):
    ph = d.get("philosophy") or {}
    pts = as_dicts(ph.get("points"), "title")[:4]
    strengths = as_dicts(d.get("strengths"))[:6]
    intro = clean(ph.get("intro") or gv(d, "identity", "intro") or "")
    if not (pts or strengths or intro):
        return None
    s = blank_slide(prs)
    y = head(s, P, "PHILOSOPHY",
             clean(ph.get("headline") or gv(d, "identity", "headline")
                   or "공부하는 방법이 결과를 만듭니다"),
             intro)

    for i, p in enumerate(pts):
        add_pill(s, ML, y, 0.46, 0.28, f"{i+1:02d}", fill=P["primary"],
                 color=P["onDark"], size=8.0, radius=0.2)
        add_text(s, ML + 0.62, y, CW - 0.62, 0.28, clean(p.get("title") or ""),
                 size=FS.body + 0.75, bold=True, color=P["text"],
                 anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.0)
        y += 0.55
    if pts:
        y += 0.25

    if strengths:
        add_text(s, ML, y, CW, 0.24, "우리 학원의 강점", size=FS.small,
                 bold=True, color=P["primary"], line_spacing=1.0)
        y += 0.42
        cols, gap = 2, 0.28
        cw = (CW - gap) / cols
        ch = 1.32
        for i, it in enumerate(strengths):
            r, c = divmod(i, cols)
            x = ML + c * (cw + gap)
            yy = y + r * (ch + 0.24)
            if yy + ch > SH - 1.1:
                break
            add_box(s, x, yy, cw, ch, fill=P["card"], line=P["line"], radius=R)
            add_text(s, x + 0.2, yy + 0.16, 0.5, 0.22, f"{i+1:02d}",
                     size=FS.small, bold=True, color=P["primary"], line_spacing=1.0)
            ttl = clean(it.get("title") or it.get("name") or "")
            add_text(s, x + 0.2, yy + 0.44, cw - 0.4, 0.28, ttl,
                     size=fit_size(ttl, FS.body + 1.5, cw - 0.4, 0.3),
                     bold=True, color=P["text"], line_spacing=1.1)
            body = clean(it.get("desc") or it.get("body") or "")
            if body:
                add_text(s, x + 0.2, yy + 0.78, cw - 0.4, 0.44, body,
                         size=fit_size(body, FS.body_sm, cw - 0.4, 0.44, min_size=8.0),
                         color=P["muted"], line_spacing=1.4)
    footer(s, P, d, page)
    return s


# ── 3. 커리큘럼 ─────────────────────────────────────────────────
def c_curriculum(prs, d, P, page):
    stages = as_dicts(d.get("curriculum"), "name")[:4]
    if not stages:
        return None
    s = blank_slide(prs)
    y = head(s, P, "CURRICULUM", "학년별 학습 설계",
             d.get("curriculumLead") or
             "현재 수준에서 시작해 다음 단계까지 이어지도록 설계합니다.")

    avail = SH - 1.25 - y
    ch = min(2.15, (avail - 0.24 * (len(stages) - 1)) / len(stages))
    for st in stages:
        add_box(s, ML, y, CW, ch, fill=P["card"], line=P["line"], radius=R)
        gname = clean(st.get("name") or st.get("grade") or "")
        gw = min(CW - 0.68, 0.5 + len(gname) * 0.135)
        add_pill(s, ML + 0.34, y + 0.22, gw, 0.29, gname,
                 fill=P["primary"], color=P["onDark"], size=8.0)
        ttl = clean(st.get("title") or "")
        add_text(s, ML + 0.34, y + 0.56, CW - 0.68, 0.32, ttl,
                 size=fit_size(ttl, FS.card_ttl, CW - 0.68, 0.34),
                 bold=True, color=P["text"], line_spacing=1.1)
        body = clean(st.get("desc") or st.get("body") or "")
        if body:
            add_text(s, ML + 0.34, y + 0.96, CW - 0.68, ch - 1.4, body,
                     size=fit_size(body, FS.body, CW - 0.68, ch - 1.4),
                     color=P["muted"], line_spacing=1.45)
        tags = [clean(t) for t in (st.get("tags") or []) if clean(t)][:3]
        if tags and ch > 1.6:
            tw = 1.15
            for j, t in enumerate(tags):
                add_pill(s, ML + 0.34 + j * (tw + 0.12), y + ch - 0.44,
                         tw, 0.28, t, fill=P["card2"], color=P["primary"],
                         size=8.0, radius=0.2)
        y += ch + 0.24
    footer(s, P, d, page)
    return s


# ── 4. 실적 + 특별프로그램 ──────────────────────────────────────
def c_results(prs, d, P, page):
    ach = as_dicts(d.get("achievements"))[:3]
    sp = as_dicts(d.get("specials"))[:4]
    if not (ach or sp):
        return None
    s = blank_slide(prs)
    y = head(s, P, "RESULTS & PROGRAMS",
             clean(d.get("achievementsHead") or "결과로 확인하는 학습"),
             d.get("achievementsLead"))

    if ach:
        gap = 0.22
        cw = (CW - gap * (len(ach) - 1)) / len(ach)
        ch = 1.85
        for i, it in enumerate(ach):
            x = ML + i * (cw + gap)
            add_box(s, x, y, cw, ch, fill=P["card"], line=P["line"], radius=R)
            add_text(s, x + 0.18, y + 0.22, cw - 0.36, 0.55,
                     clean(it.get("value") or it.get("num") or ""),
                     size=23.25, bold=True, color=P["primary"], line_spacing=1.0)
            add_text(s, x + 0.18, y + 0.82, cw - 0.36, 0.26,
                     clean(it.get("title") or it.get("label") or ""),
                     size=FS.body + 0.75, bold=True, color=P["text"], line_spacing=1.0)
            body = clean(it.get("desc") or "")
            if body:
                add_text(s, x + 0.18, y + 1.14, cw - 0.36, 0.6, body,
                         size=fit_size(body, FS.small, cw - 0.36, 0.6, min_size=7.5),
                         color=P["muted"], line_spacing=1.4)
        y += ch + 0.5

    if sp:
        add_text(s, ML, y, CW, 0.24, "특별 프로그램", size=FS.small,
                 bold=True, color=P["primary"], line_spacing=1.0)
        y += 0.42
        for i, it in enumerate(sp):
            if y + 0.9 > SH - 1.1:
                break
            add_text(s, ML, y, 0.42, 0.28, f"{i+1:02d}", size=FS.body + 1.5,
                     bold=True, color=P["accent"], line_spacing=1.0)
            add_text(s, ML + 0.55, y, CW - 0.55, 0.28,
                     clean(it.get("title") or it.get("name") or ""),
                     size=FS.body + 1.5, bold=True, color=P["text"],
                     anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.0)
            body = clean(it.get("desc") or it.get("body") or "")
            if body:
                add_text(s, ML + 0.55, y + 0.34, CW - 0.55, 0.5,
                         wrap_words(body, 46)[:2], size=FS.body_sm,
                         color=P["muted"], line_spacing=1.45)
            y += 1.02
    footer(s, P, d, page)
    return s


# ── 5. 시간표 + 학습관리 ────────────────────────────────────────
def c_schedule(prs, d, P, page):
    groups = [g for g in as_dicts(d.get("timetables"), "group") if g.get("rows")]
    mg = as_dicts(d.get("management"))[:4]
    if not (groups or mg):
        return None
    s = blank_slide(prs)
    y = head(s, P, "SCHEDULE & CARE", "수업 시간표와 학습 관리",
             "시간표는 학기별로 조정될 수 있으니 상담 시 확인해 주세요."
             if groups else None)

    if groups:
        cols = 2 if len(groups) > 1 else 1
        gap = 0.3
        tw = (CW - gap * (cols - 1)) / cols
        rows_n = -(-min(len(groups), 4) // cols)
        block = min(2.9, (5.3 - 0.3 * (rows_n - 1)) / rows_n)

        for i, g in enumerate(groups[:4]):
            r, c = divmod(i, cols)
            x = ML + c * (tw + gap)
            yy = y + r * (block + 0.3)

            rws = as_dicts(g.get("rows"), "name")
            total = len(rws)
            maxr = max(1, int((block - 0.38 - 0.42) / 0.32))
            shown = rws[:maxr]
            if not shown:
                continue
            omitted = total - len(shown)

            gname = clean(g.get("group") or "")
            if omitted > 0:
                gname = f"{gname}  (총 {total}개 중 {len(shown)}개)"
            add_text(s, x, yy, tw, 0.26, gname, size=FS.body + 0.75,
                     bold=True, color=P["text"], line_spacing=1.0)

            rws = shown
            row_h = min(0.36, max(0.28,
                                  (block - 0.38 - 0.42) / max(len(rws), 1)))
            gf = s.shapes.add_table(len(rws) + 1, 2, Inches(x),
                                    Inches(yy + 0.36), Inches(tw),
                                    Inches(0.42 + row_h * len(rws)))
            tbl = gf.table
            c1 = tw * 0.56
            _style_table(tbl, P, [c1, tw - c1], 0.42, row_h)
            _set_cell(tbl.cell(0, 0), "반", P, 8.0, True, P["onDark"])
            _set_cell(tbl.cell(0, 1), "요일 · 시간", P, 8.0, True,
                      P["onDark"], PP_ALIGN.RIGHT)
            for ri, row in enumerate(rws, start=1):
                nm = clean(row.get("name") or row.get("class") or "")
                tm = clean(row.get("time") or row.get("when") or "")
                _set_cell(tbl.cell(ri, 0), nm, P,
                          fit_size(nm, FS.body_sm, c1 - 0.24, row_h,
                                   min_size=7.0),
                          True, P["text"])
                _set_cell(tbl.cell(ri, 1), tm, P, 8.0, False,
                          P["muted"], PP_ALIGN.RIGHT)
        y += rows_n * (block + 0.3) + 0.2

    if mg and y + 1.2 < SH - 1.1:
        add_text(s, ML, y, CW, 0.24, "학습 관리", size=FS.small,
                 bold=True, color=P["primary"], line_spacing=1.0)
        y += 0.4
        cw2 = (CW - 0.24) / 2
        for i, it in enumerate(mg):
            r, c = divmod(i, 2)
            x = ML + c * (cw2 + 0.24)
            yy = y + r * 1.0
            if yy + 0.88 > SH - 1.05:
                break
            add_box(s, x, yy, cw2, 0.88, fill=P["card"], line=P["line"], radius=R)
            key = clean(it.get("key") or it.get("title") or "")
            add_pill(s, x + 0.16, yy + 0.14, 0.72, 0.26, key, fill=P["primary"],
                     color=P["onDark"], size=8.0, radius=0.2)
            body = clean(it.get("desc") or it.get("body") or "")
            add_text(s, x + 0.16, yy + 0.47, cw2 - 0.32, 0.34,
                     wrap_words(body, 30)[:2], size=8.0,
                     color=P["muted"], line_spacing=1.35)
    footer(s, P, d, page)
    return s


# ── 6. 입학 + FAQ + 연락처 ──────────────────────────────────────
def c_admission(prs, d, P, page):
    steps = as_dicts(d.get("admission"))[:4]
    faq = [x for x in as_dicts(d.get("faq"), "q")
           if clean(x.get("q") or x.get("question"))][:3]
    b = d.get("basic", {})
    if not (steps or faq):
        return None
    s = blank_slide(prs)
    y = head(s, P, "ADMISSION", "등록 안내",
             clean(d.get("admissionNote") or ""))

    if steps:
        gap = 0.2
        cw = (CW - gap * (len(steps) - 1)) / len(steps)
        ch = 1.55
        for i, st in enumerate(steps):
            x = ML + i * (cw + gap)
            add_box(s, x, y, cw, ch, fill=P["card"], line=P["line"], radius=R)
            add_text(s, x + 0.16, y + 0.18, 0.4, 0.32, str(i + 1),
                     size=18, bold=True, color=P["accent"], line_spacing=1.0)
            ttl = clean(st.get("title") or st.get("name") or "")
            add_text(s, x + 0.16, y + 0.58, cw - 0.32, 0.28, ttl,
                     size=fit_size(ttl, FS.body + 0.75, cw - 0.32, 0.3),
                     bold=True, color=P["text"], line_spacing=1.1)
            body = clean(st.get("desc") or "")
            if body:
                add_text(s, x + 0.16, y + 0.9, cw - 0.32, 0.5, body,
                         size=fit_size(body, 8.0, cw - 0.32, 0.5, min_size=7.0),
                         color=P["muted"], line_spacing=1.35)
        y += ch + 0.5

    if faq:
        add_text(s, ML, y, CW, 0.24, "자주 묻는 질문", size=FS.small,
                 bold=True, color=P["primary"], line_spacing=1.0)
        y += 0.4
        for it in faq:
            if y + 0.9 > SH - 2.6:
                break
            q = clean(it.get("q") or it.get("question") or "")
            a = clean(it.get("a") or it.get("answer") or "")
            add_text(s, ML, y, CW, 0.26, "Q. " + q, size=FS.body,
                     bold=True, color=P["text"], line_spacing=1.2)
            add_text(s, ML, y + 0.3, CW, 0.5, "A. " + a, size=FS.body_sm,
                     color=P["muted"], line_spacing=1.45)
            y += 1.0

    # 연락처 블록
    by = SH - 2.35
    add_box(s, ML, by, CW, 1.4, fill=P["deep"], radius=R)
    add_text(s, ML + 0.32, by + 0.24, CW - 0.64, 0.26, "상담 · 예약",
             size=FS.small, bold=True, color=P["accent"], line_spacing=1.0)
    add_text(s, ML + 0.32, by + 0.58, CW * 0.5, 0.36, clean(b.get("phone") or ""),
             size=19.5, bold=True, color=P["onDark"], line_spacing=1.0)
    addr = clean(b.get("address") or "")
    if addr:
        add_text(s, ML + 0.32, by + 1.02, CW - 0.64, 0.26,
                 wrap_words(addr, 44)[:1], size=FS.small,
                 color=P["onDarkSub"], line_spacing=1.0)
    footer(s, P, d, page)
    return s


def build(data, out=None, palette=None):
    P = get_palette(palette)
    d = data or {}
    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)

    c_cover(prs, d, P)
    page = 2
    for fn in (c_philosophy, c_curriculum, c_results, c_schedule, c_admission):
        if fn(prs, d, P, page) is not None:
            page += 1

    if out is None:
        import io
        bio = io.BytesIO()
        prs.save(bio)
        bio.seek(0)
        return bio
    prs.save(out)
    return out


if __name__ == "__main__":
    import json, sys
    src = sys.argv[1] if len(sys.argv) > 1 else "sample.json"
    pal = sys.argv[2] if len(sys.argv) > 2 else "navy_gold"
    with open(src, encoding="utf-8") as f:
        data = json.load(f)
    print(build(data, f"navy_catalog_{pal}.pptx", pal))
